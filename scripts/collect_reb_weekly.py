#!/usr/bin/env python3
"""Collect the exact official REB weekly apartment sale/rent time series.

Source: 한국부동산원 R-ONE 공개자료실 > 주간아파트가격동향조사 시계열통계표.
The workbook is downloaded from the public attachment UI and parsed without a private API.
"""
from __future__ import annotations
import datetime,json,pathlib,re,tempfile,zipfile
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,TimeoutError as PlaywrightTimeoutError

ROOT=pathlib.Path(__file__).resolve().parents[1]
OUT=ROOT/'data_sources/reb_weekly.json'
LIST_URL='https://www.reb.or.kr/r-one/portal/bbs/statdata/searchBulletinPage.do'
TARGET_TITLE='주간아파트가격동향조사'

def norm_date(v):
    if isinstance(v,(datetime.datetime,datetime.date)):
        return v.strftime('%Y%m%d')
    s=''.join(re.findall(r'\d',str(v or '')))
    if len(s)>=8:
        try:return datetime.datetime.strptime(s[:8],'%Y%m%d').strftime('%Y%m%d')
        except ValueError:return None
    if len(s)==6:
        # workbook can encode dates as YYMMDD in older rows
        try:return datetime.datetime.strptime('20'+s,'%Y%m%d').strftime('%Y%m%d')
        except ValueError:return None
    return None

def number(v):
    if v is None:return None
    if isinstance(v,(int,float)):return float(v)
    s=str(v).strip().replace(',','')
    if not s or s in ('-','—'):return None
    try:return float(s)
    except ValueError:return None

def download_workbook(dest:pathlib.Path):
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=True)
        ctx=browser.new_context(accept_downloads=True,locale='ko-KR')
        page=ctx.new_page()
        page.goto(LIST_URL,wait_until='domcontentloaded',timeout=45000)
        page.wait_for_timeout(5000)
        rows=page.locator('table tbody tr')
        if rows.count()==0:
            print(json.dumps({'debug':'reb_board_empty','url':page.url,'title':page.title(),'body':page.locator('body').inner_text()[:5000]},ensure_ascii=False))
        hit=None
        for i in range(rows.count()):
            row=rows.nth(i)
            txt=row.inner_text().strip()
            if TARGET_TITLE in txt:
                hit=row
                break
        if hit is None:
            texts=[rows.nth(i).inner_text().strip() for i in range(min(rows.count(),30))]
            print(json.dumps({'debug':'reb_board_rows','url':page.url,'rows':texts},ensure_ascii=False))
            raise RuntimeError('REB weekly time-series bulletin row not found')
        # Open attachment layer from the title cell/link.
        links=hit.locator('a')
        if links.count():
            links.first.evaluate("el => el.click()")
        else:
            tds=hit.locator('td')
            if tds.count()<3:raise RuntimeError('REB bulletin row malformed')
            tds.nth(2).click()
        page.wait_for_timeout(900)
        frame=None
        for _ in range(20):
            frame=page.frame(name='raonkuploader_frame_kupload')
            if frame:break
            page.wait_for_timeout(250)
        if frame is not None:
            files=frame.locator('#file_list > li')
            if files.count()<1:raise RuntimeError('REB attachment list empty')
            chosen=None;file_names=[]
            for i in range(files.count()):
                li=files.nth(i);txt=li.inner_text().strip();file_names.append(txt);low=txt.lower()
                if '.xlsx' in low or '.xls' in low or '시계열' in txt:
                    chosen=li
                    if '.xlsx' in low:break
            if chosen is None:chosen=files.nth(1 if files.count()>1 else 0)
            try:chosen.locator('ul li').nth(1).click()
            except Exception:chosen.click()
            frame.locator('#button_download').wait_for(state='visible',timeout=10000)
            try:
                with page.expect_download(timeout=20000) as di:frame.locator('#button_download').click()
                dl=di.value
            except PlaywrightTimeoutError:
                raise RuntimeError('REB workbook download did not start; files='+json.dumps(file_names,ensure_ascii=False))
            raw_path=dest.with_suffix(pathlib.Path(dl.suggested_filename).suffix or '.bin')
            dl.save_as(str(raw_path))
        else:
            att=page.locator('#notice-attach-sect a.atchFile')
            if att.count()<1:
                raise RuntimeError('REB current attachment link not found')
            attachment_text=att.first.inner_text().strip()
            try:
                with page.expect_download(timeout=30000) as di:
                    att.first.evaluate("el => el.click()")
                dl=di.value
            except PlaywrightTimeoutError:
                raise RuntimeError('REB current attachment download did not start: '+attachment_text)
            raw_path=dest.with_suffix(pathlib.Path(dl.suggested_filename).suffix or '.zip')
            dl.save_as(str(raw_path))
        if raw_path.suffix.lower()=='.zip':
            with zipfile.ZipFile(raw_path) as z:
                names=z.namelist()
                books=[n for n in names if n.lower().endswith('.xlsx')]
                if not books:
                    raise RuntimeError('REB zip has no xlsx workbook: '+json.dumps(names,ensure_ascii=False))
                sale_entry=next((n for n in books if '매매가격지수/' in n and '(주) 매매가격지수.xlsx' in n),None)
                rent_entry=next((n for n in books if '전세가격지수/' in n and '(주) 전세가격지수.xlsx' in n),None)
                if not sale_entry or not rent_entry:
                    raise RuntimeError('REB exact sale/rent workbooks missing: '+json.dumps(books,ensure_ascii=False))
                rent_dest=dest.with_name(dest.stem+'_rent.xlsx')
                dest.write_bytes(z.read(sale_entry))
                rent_dest.write_bytes(z.read(rent_entry))
                print(json.dumps({'downloaded_zip':raw_path.name,'sale_entry':sale_entry,'rent_entry':rent_entry,'entries':names[:40]},ensure_ascii=False))
        else:
            if raw_path!=dest:dest.write_bytes(raw_path.read_bytes())
        browser.close()

def sheet_for(wb,kind):
    exact='매매지수' if kind=='sale' else '전세지수'
    if exact in wb.sheetnames:return wb[exact]
    needle='매매' if kind=='sale' else '전세'
    for name in wb.sheetnames:
        if needle in name and '지수' in name and '변동' not in name:
            return wb[name]
    raise RuntimeError(f'{kind} index sheet missing: {wb.sheetnames}')

def find_seoul_col(ws):
    candidates=[]
    for c in range(1,ws.max_column+1):
        hit=False
        for r in range(1,min(ws.max_row,60)+1):
            v=ws.cell(r,c).value
            if v is None:continue
            s=str(v).strip().replace(' ','')
            if s=='서울' or s=='서울특별시':
                hit=True;break
        if hit:candidates.append(c)
    if not candidates:
        sample=[]
        for r in range(1,min(ws.max_row,25)+1):
            row=[ws.cell(r,c).value for c in range(1,min(ws.max_column,25)+1)]
            sample.append(row)
        print(json.dumps({'debug':'reb_sheet_header','sheet':ws.title,'max_row':ws.max_row,'max_col':ws.max_column,'rows':sample},ensure_ascii=False,default=str))
        raise RuntimeError(f'Seoul column not found in {ws.title}')
    return min(candidates)

def find_date_col(ws):
    scored=[]
    for c in range(1,min(ws.max_column,12)+1):
        n=0
        for r in range(1,ws.max_row+1):
            if norm_date(ws.cell(r,c).value):n+=1
        scored.append((n,c))
    n,c=max(scored)
    if n<50:raise RuntimeError(f'date column not found in {ws.title}: {scored}')
    return c

def extract(ws):
    c=find_seoul_col(ws);dc=find_date_col(ws)
    out=[]
    for r in range(1,ws.max_row+1):
        d=norm_date(ws.cell(r,dc).value)
        v=number(ws.cell(r,c).value)
        if d and v is not None and 40<=v<=180:
            out.append({'date':d,'value':v})
    by={x['date']:x for x in out}
    out=[by[k] for k in sorted(by)]
    if len(out)<100:
        raise RuntimeError(f'{ws.title} Seoul history too short: {len(out)}')
    return out,c

def pct(a,b):
    return round((b/a-1)*100,2) if a not in (None,0) and b is not None else None

def lag(series,n):
    return pct(series[-1-n]['value'],series[-1]['value']) if len(series)>n else None

def main():
    with tempfile.TemporaryDirectory() as td:
        book=pathlib.Path(td)/'reb_weekly.xlsx'
        download_workbook(book)
        rent_book=book.with_name(book.stem+'_rent.xlsx')
        for label,pth in [('sale',book),('rent',rent_book)]:
            with zipfile.ZipFile(pth) as zz:
                names=zz.namelist()
                sizes={n:zz.getinfo(n).file_size for n in names}
                inspect=[n for n in names if n.startswith(('xl/worksheets/','xl/sharedStrings','xl/externalLinks/','xl/queryTables/','xl/connections'))]
                snippets={}
                for n in inspect[:20]:
                    try:snippets[n]=zz.read(n)[:1200].decode('utf-8',errors='replace')
                    except Exception:pass
                print(json.dumps({'debug':'reb_xlsx_structure','kind':label,'names':names,'sizes':sizes,'snippets':snippets},ensure_ascii=False))
        # REB workbooks declare the incorrect dimension A1 even though sheet XML contains
        # the full table. Normal mode parses actual cell records instead of trusting that dimension.
        swb=load_workbook(book,read_only=False,data_only=True)
        rwb=load_workbook(rent_book,read_only=False,data_only=True)
        sws=swb[swb.sheetnames[0]];rws=rwb[rwb.sheetnames[0]]
        sale,scol=extract(sws);rent,rcol=extract(rws)
    sm={x['date']:x for x in sale};rm={x['date']:x for x in rent}
    common=sorted(set(sm)&set(rm))
    if len(common)<100:raise RuntimeError(f'common REB weekly history too short: {len(common)}')
    sale=[sm[d] for d in common];rent=[rm[d] for d in common]
    latest=common[-1];start='20250203'
    if start not in common:
        later=[d for d in common if d>=start]
        if not later:raise RuntimeError('2025-02-03 start unavailable')
        start=later[0]
        gap=(datetime.datetime.strptime(start,'%Y%m%d')-datetime.datetime.strptime('20250203','%Y%m%d')).days
        if gap>14:raise RuntimeError(f'2025-02-03 start gap too large: {start}')
    i=common.index(start)
    acc={'requested_start_date':'20250203','start_date':start,'end_date':latest,
         'interval_weeks':round((datetime.datetime.strptime(latest,'%Y%m%d')-datetime.datetime.strptime(start,'%Y%m%d')).days/7,1),
         'sale_pct':pct(sale[i]['value'],sale[-1]['value']),'rent_pct':pct(rent[i]['value'],rent[-1]['value'])}
    if acc['sale_pct'] is not None and acc['rent_pct'] is not None:
        acc['sale_minus_rent_pp']=round(acc['sale_pct']-acc['rent_pct'],2)
    out={'status':'connected','source':'한국부동산원 R-ONE 주간아파트가격동향조사 시계열통계표',
         'scope':{'region':'서울','housing_type':'아파트'},'frequency':'weekly',
         'latest':{'date':latest,'sale_index':sale[-1]['value'],'rent_index':rent[-1]['value']},
         'momentum':{'sale_4w_pct':lag(sale,4),'sale_13w_pct':lag(sale,13),
                     'rent_4w_pct':lag(rent,4),'rent_13w_pct':lag(rent,13)},
         'accumulation':acc,'series':{'sale':sale,'rent':rent},
         'workbook':{'bulletin_title':TARGET_TITLE,'sale_sheet':sws.title,'rent_sheet':rws.title,
                     'sale_seoul_column':scol,'rent_seoul_column':rcol},
         'role':'crosscheck_context_only','production_model_weight_changed':False,
         'collection_method':'public_reb_timeseries_workbook',
         'collected_at':datetime.datetime.now(datetime.timezone.utc).isoformat()}
    OUT.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'status':'ok','latest':out['latest'],'momentum':out['momentum'],'accumulation':acc,'workbook':out['workbook']},ensure_ascii=False))

if __name__=='__main__':main()
